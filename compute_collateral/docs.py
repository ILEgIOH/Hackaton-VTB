"""Заключение об оценке: HTML (просмотр), PDF (reportlab) и Word (python-docx). Разделы включаются галочками (экран D02)."""
from __future__ import annotations

import html as _html
import io
from datetime import date
from pathlib import Path

from . import engine

FONT_DIR = Path(__file__).resolve().parent / "fonts"
SECTIONS = [("valuation", "Стоимость и три метода оценки"), ("forecast", "Прогноз и график покрытия"),
            ("equipment", "Список оборудования и состояние"), ("schedule", "График погашения")]
SECTION_KEYS = [k for k, _ in SECTIONS]
SCEN = dict(opt="Оптимистичный", base="Базовый", pess="Пессимистичный")


def rub(v, d=1):
    return f"{v / 1e6:,.{d}f} млн ₽".replace(",", " ").replace(".", ",")


def c2(v):
    return "—" if v is None else f"{v:.2f}".replace(".", ",")


def pct(v, d=1):
    return f"{v * 100:.{d}f} %".replace(".", ",")


def dmy(iso):
    return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}"


def build(deal: dict, detail: dict, servers: list[dict], ai: dict | None, author: str, sections=None) -> dict:
    """Формат-независимая модель документа: заголовок, ключевые цифры, абзацы и таблицы."""
    sections = [s for s in (sections or SECTION_KEYS) if s in SECTION_KEYS]
    v, fc, rec, eq = detail["valuation"], detail["forecast"], detail["recommendation"], detail["equipment"]
    doc = dict(title="Заключение об оценке залога", subtitle=f"{deal['borrower']} · сделка {deal.get('contract', '')} · {deal['location']}",
               meta=f"Подготовил: {author} · {date.today().strftime('%d.%m.%Y')}", blocks=[], sections=sections)
    B = doc["blocks"]
    B.append(("kpis", [("Рыночная стоимость", rub(v["market"])), ("Ликвидационная стоимость", rub(v["liquidation"])),
                       ("Рекомендуемый LTV", pct(rec["ltv"]))]))
    B.append(("h", "Параметры кредита"))
    B.append(("p", f"Сумма {rub(fc['amount'])}, срок {fc['term']} мес., ставка {str(fc['rate']).replace('.', ',')} % годовых, "
                   f"ковенант покрытия {c2(fc['covenant'])}. Текущее покрытие: {c2(detail['coverage_now'])}. "
                   f"Рекомендуемая максимальная сумма: {rub(rec['max_loan'])}."))
    if "valuation" in sections:
        B.append(("h", "Из чего сложилась оценка"))
        B.append(("table", ["Подход", "Вес", "Стоимость"], [
            ["Затратный", "30 %", rub(v["cost"])], ["Сравнительный", "40 %", rub(v["comparative"])], ["Доходный", "30 %", rub(v["income"])],
            ["Среднее с учётом весов", "", rub(v["weighted"])], ["Поправка за износ", "", "−" + pct(v["wear"])],
            ["Рыночная стоимость", "", rub(v["market"])]]))
        if v["low_data"]:
            B.append(("p", "Мало рыночных цен по моделям: " + ", ".join(v["low_data"]) + ". Оценка по мировым ценам, точность ±15 %."))
    if "forecast" in sections:
        B.append(("h", "Прогноз покрытия"))
        B.append(("chart", fc["points"]))
        B.append(("table", ["Дата", "Остаток долга", "Базовый", "Пессим."],
                  [[dmy(r["date"]), rub(r["debt"]), "погашен" if r["paid"] else c2(r["base"]), "погашен" if r["paid"] else c2(r["pess"])]
                   for r in detail["half_year"]]))
    if "equipment" in sections:
        B.append(("h", "Оборудование"))
        B.append(("p", f"{v['cards']} карт. Серийные номера подтверждены: {eq['matched']} из {eq['total']}. "
                       f"Телеметрия: {({'agent': 'агент на серверах', 'file': 'файл nvidia-smi', 'none': 'нет (поправка 7 %)'}).get(deal.get('telemetry_mode'), '—')}."))
        B.append(("table", ["Сервер", "Карт", "Наработка, ч", "Загрузка", "Темп.", "Серийные номера", "Связь"],
                  [[s["name"], str(s["cards"]), "—" if s["hours"] is None else f"{s['hours']:,}".replace(",", " "),
                    "—" if s["util"] is None else f"{s['util']} %", "—" if s["temp"] is None else f"{s['temp']} °C",
                    f"{s['matched']} из {s['cards']}" if s["seen"] else "не проверены",
                    {True: "на связи", False: "нет связи", None: "—"}[s["online"]]] for s in servers]))
    if "schedule" in sections:
        B.append(("h", "График погашения"))
        sched = engine.schedule(fc["amount"], fc["term"], fc["rate"], deal["loan"].get("schedule", "annuity"))
        start = engine.add_months(date.fromisoformat(deal["loan"]["issue_date"]), 0)
        B.append(("table", ["№", "Дата", "Остаток долга после платежа"],
                  [[str(k), engine.add_months(start, k).strftime("%d.%m.%Y"), rub(sched[k], 2)] for k in range(1, len(sched))]))
    B.append(("h", "Риски и замечания"))
    risks = []
    if detail["gap"]:
        g = detail["gap"]
        risks.append(f"В пессимистичном сценарии залог станет дешевле долга {dmy(g['first_date'])}; разрыв продлится {g['months']} мес., максимум {rub(g['max_gap'])}.")
    if rec["exceeds"]:
        risks.append(f"Запрошенная сумма {rub(fc['amount'])} выше рекомендуемой {rub(rec['max_loan'])} (LTV {pct(rec['ltv'])}).")
    for m in eq["mismatches"]:
        risks.append(f"Серийный номер не совпадает: {m['server']}, GPU {m['gpu']} — по договору {m['contract']}, установлен {m['actual']}.")
    for s in eq.get("silent", []):
        risks.append(f"Сервер {s['server']} не передаёт данные, его карты оценены с поправкой 7 %.")
    if ai:
        risks.append("Комментарий ИИ-наблюдателя (справочно, в расчёте не участвует): " + ai["summary"])
    B.append(("list", risks or ["Существенных рисков не выявлено."]))
    return doc


# ---------- HTML ----------

def to_html(doc: dict) -> str:
    e = _html.escape
    out = [f"""<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>{e(doc['title'])} — {e(doc['subtitle'])}</title>
<style>body{{font:14px/1.5 Manrope,Arial,sans-serif;color:#051F45;max-width:820px;margin:32px auto;padding:0 20px}}h1{{font-size:24px;margin:0}}
h2{{font-size:16px;margin:26px 0 8px;border-bottom:1px solid #EBDCC5;padding-bottom:4px}}table{{border-collapse:collapse;width:100%;font-size:13px}}
td,th{{padding:6px 8px;border-bottom:1px solid #EBDCC5;text-align:left}}th{{color:#3E4E68}}.big{{font-size:26px;font-weight:700}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:16px}}.m{{color:#3E4E68}}svg{{width:100%;height:auto}}</style></head><body>
<p class="m">ComputeCollateral</p><h1>{e(doc['title'])}</h1><p class="m">{e(doc['subtitle'])}<br>{e(doc['meta'])}</p>"""]
    for b in doc["blocks"]:
        kind = b[0]
        if kind == "kpis":
            out.append('<div class="grid">' + "".join(f'<div><div class="m">{e(k)}</div><div class="big">{e(v)}</div></div>' for k, v in b[1]) + "</div>")
        elif kind == "h":
            out.append(f"<h2>{e(b[1])}</h2>")
        elif kind == "p":
            out.append(f"<p>{e(b[1])}</p>")
        elif kind == "list":
            out.append("<ul>" + "".join(f"<li>{e(x)}</li>" for x in b[1]) + "</ul>")
        elif kind == "table":
            out.append("<table><tr>" + "".join(f"<th>{e(h)}</th>" for h in b[1]) + "</tr>" +
                       "".join("<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in b[2]) + "</table>")
        elif kind == "chart":
            out.append(_svg_chart(b[1]))
    out.append("</body></html>")
    return "".join(out)


def _svg_chart(pts):
    W, H, ml, mt, mb = 760, 260, 40, 10, 24
    mx = max(max(p["opt"], p["debt"]) for p in pts) * 1.05
    x = lambda i: ml + (W - ml - 10) * i / max(1, len(pts) - 1)
    y = lambda v: mt + (H - mt - mb) * (1 - v / mx)
    line = lambda k: "M" + " L".join(f"{x(i):.1f} {y(p[k]):.1f}" for i, p in enumerate(pts))
    debt = "M" + " ".join((f"{x(i):.1f} {y(p['debt']):.1f}" if i == 0 else f"H{x(i):.1f} V{y(p['debt']):.1f}") for i, p in enumerate(pts))
    return (f'<svg viewBox="0 0 {W} {H}"><rect width="{W}" height="{H}" fill="#fff"/>'
            f'<path d="{line("opt")}" fill="none" stroke="#9fc9d4" stroke-width="2"/><path d="{line("base")}" fill="none" stroke="#4C9DB0" stroke-width="3"/>'
            f'<path d="{line("pess")}" fill="none" stroke="#7b8799" stroke-width="2" stroke-dasharray="5 4"/><path d="{debt}" fill="none" stroke="#051F45" stroke-width="2.5"/>'
            f'<text x="{ml}" y="{H - 6}" font-size="11" fill="#3E4E68">{dmy(pts[0]["date"])}</text><text x="{W - 60}" y="{H - 6}" font-size="11" fill="#3E4E68">{dmy(pts[-1]["date"])}</text>'
            '<text x="50" y="22" font-size="11" fill="#3E4E68">светлая — оптимистичный · бирюзовая — базовый · пунктир — пессимистичный · тёмная ступенчатая — остаток долга</text></svg>')


# ---------- Word ----------

def to_docx(doc: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10)
    d.add_heading(doc["title"], 0)
    p = d.add_paragraph(doc["subtitle"])
    p.add_run("\n" + doc["meta"]).font.color.rgb = RGBColor(0x3E, 0x4E, 0x68)
    for b in doc["blocks"]:
        kind = b[0]
        if kind == "kpis":
            t = d.add_table(rows=2, cols=len(b[1]))
            for i, (k, v) in enumerate(b[1]):
                t.cell(0, i).text = k
                t.cell(1, i).text = v
                t.cell(1, i).paragraphs[0].runs[0].bold = True
        elif kind == "h":
            d.add_heading(b[1], 2)
        elif kind == "p":
            d.add_paragraph(b[1])
        elif kind == "list":
            for x in b[1]:
                d.add_paragraph(x, style="List Bullet")
        elif kind == "table":
            t = d.add_table(rows=1, cols=len(b[1]))
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(b[1]):
                t.rows[0].cells[i].text = h
            for r in b[2]:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = c
        elif kind == "chart":
            d.add_paragraph("График покрытия приведён в PDF-версии заключения и в приложении.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------- PDF ----------

def to_pdf(doc: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, Preformatted
    from reportlab.graphics.shapes import Drawing, Line, PolyLine, String

    if "CCSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("CCSans", str(FONT_DIR / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("CCSans-B", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("CCSans", normal="CCSans", bold="CCSans-B")
    navy, slate = colors.HexColor("#051F45"), colors.HexColor("#3E4E68")
    base = ParagraphStyle("b", fontName="CCSans", fontSize=9, leading=13, textColor=navy)
    h1 = ParagraphStyle("h1", parent=base, fontName="CCSans-B", fontSize=18, leading=22, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=base, fontName="CCSans-B", fontSize=11, leading=14, spaceBefore=12, spaceAfter=4)
    muted = ParagraphStyle("m", parent=base, textColor=slate)
    cell = ParagraphStyle("c", parent=base, fontSize=8, leading=10)
    esc = lambda s: _html.escape(str(s))
    story = [Paragraph(esc(doc["title"]), h1), Paragraph(esc(doc["subtitle"]), muted), Paragraph(esc(doc["meta"]), muted), Spacer(1, 6)]
    for b in doc["blocks"]:
        kind = b[0]
        if kind == "kpis":
            t = Table([[Paragraph(esc(k), muted) for k, _ in b[1]], [Paragraph(f"<b>{esc(v)}</b>", ParagraphStyle("k", parent=base, fontName="CCSans-B", fontSize=14, leading=18)) for _, v in b[1]]])
            story.append(t)
        elif kind == "h":
            story.append(Paragraph(esc(b[1]), h2))
        elif kind == "p":
            story.append(Paragraph(esc(b[1]), base))
        elif kind == "list":
            for x in b[1]:
                story.append(Paragraph("• " + esc(x), base))
        elif kind == "table":
            data = [[Paragraph(f"<b>{esc(h)}</b>", cell) for h in b[1]]] + [[Paragraph(esc(c), cell) for c in r] for r in b[2]]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#EBDCC5")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FBF3E4")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            story.append(t)
        elif kind == "chart":
            pts = b[1]
            W, H = 170 * mm, 55 * mm
            dr = Drawing(W, H)
            mx = max(max(p["opt"], p["debt"]) for p in pts) * 1.05
            x = lambda i: 8 * mm + (W - 12 * mm) * i / max(1, len(pts) - 1)
            y = lambda v: 6 * mm + (H - 10 * mm) * v / mx
            dr.add(Line(8 * mm, 6 * mm, W - 4 * mm, 6 * mm, strokeColor=colors.HexColor("#EBDCC5")))
            for key, col, w, dash in (("opt", "#9fc9d4", 1, None), ("base", "#4C9DB0", 1.8, None), ("pess", "#7b8799", 1.2, [3, 2])):
                dr.add(PolyLine([c for i, p in enumerate(pts) for c in (x(i), y(p[key]))], strokeColor=colors.HexColor(col), strokeWidth=w, strokeDashArray=dash))
            step = []
            for i, p in enumerate(pts):
                if i:
                    step += [x(i), y(pts[i - 1]["debt"])]
                step += [x(i), y(p["debt"])]
            dr.add(PolyLine(step, strokeColor=navy, strokeWidth=1.6))
            dr.add(String(8 * mm, 1 * mm, dmy(pts[0]["date"]), fontName="CCSans", fontSize=7, fillColor=slate))
            dr.add(String(W - 22 * mm, 1 * mm, dmy(pts[-1]["date"]), fontName="CCSans", fontSize=7, fillColor=slate))
            dr.add(String(10 * mm, H - 4 * mm, "оптимистичный · базовый · пессимистичный (пунктир) · остаток долга (тёмная линия), млн ₽: до " + f"{mx / 1e6:.0f}",
                          fontName="CCSans", fontSize=6.5, fillColor=slate))
            story.append(dr)
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                      title=doc["title"], author="ComputeCollateral").build(story)
    return buf.getvalue()


def render(doc: dict, fmt: str) -> tuple[bytes, str, str]:
    """-> (байты, MIME, расширение). Если библиотеки нет, вызывающий получает ImportError."""
    if fmt == "pdf":
        return to_pdf(doc), "application/pdf", "pdf"
    if fmt == "docx":
        return to_docx(doc), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    return to_html(doc).encode("utf-8"), "text/html; charset=utf-8", "html"
