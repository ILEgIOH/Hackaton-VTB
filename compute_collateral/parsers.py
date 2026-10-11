"""Разбор файлов: перечень оборудования заёмщика (Excel/CSV) и снимок nvidia-smi."""
from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from datetime import date, datetime

from .market import normalize_model

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 1000
SERIAL_RE = re.compile(r"^\d{13}$")

TEMPLATE_CSV = (
    "Сервер;Модель;Серийный номер;Дата ввода\r\n"
    "srv-01;H100 SXM 80 ГБ;1654023118421;12.02.2025\r\n"
    "srv-01;H100 SXM 80 ГБ;1654023118422;12.02.2025\r\n"
    "srv-02;H100 SXM 80 ГБ;1654023118908;03.06.2025\r\n"
)

HEAD_SYNONYMS = {
    "server": ("сервер", "server", "host", "hostname", "узел"),
    "model": ("модель", "model", "gpu", "name", "карта"),
    "serial": ("серийный", "serial", "s/n", "sn"),
    "date": ("дата", "date", "commission", "ввод"),
}
COL_TITLE = {"server": "Сервер", "model": "Модель", "serial": "Серийный номер", "date": "Дата ввода"}


class FileError(ValueError):
    pass


def _rows(filename: str, data: bytes) -> list[list]:
    if not data:
        raise FileError("Файл пуст")
    if len(data) > MAX_FILE_BYTES:
        raise FileError("Файл больше 10 МБ")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return [list(r) for r in wb.active.iter_rows(values_only=True)]
        except Exception:
            raise FileError("Не удалось прочитать Excel-файл. Сохраните его как .xlsx или .csv") from None
    if not name.endswith((".csv", ".txt", ".tsv")):
        raise FileError("Нужен файл Excel (.xlsx) или CSV")
    text = None
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FileError("Не удалось прочитать файл: неизвестная кодировка")
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _col_map(header: list) -> dict:
    cols = {}
    for i, h in enumerate(header):
        hl = str(h or "").strip().lower()
        for key, syns in HEAD_SYNONYMS.items():
            if key not in cols and any(s in hl for s in syns):
                cols[key] = i
    return cols


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def parse_equipment(filename: str, data: bytes, today: date | None = None) -> dict:
    """Возвращает построчную проверку (экран N04/N05) и серверы только из корректных строк."""
    today = today or date.today()
    raw = [(i + 1, r) for i, r in enumerate(_rows(filename, data)) if any(str(c or "").strip() for c in r)]
    if not raw:
        raise FileError("Файл пуст")
    cols = _col_map(raw[0][1])
    missing = [COL_TITLE[k] for k in ("server", "model", "serial", "date") if k not in cols]
    if missing:
        raise FileError("В файле нет колонки «" + "», «".join(missing) + "». Скачайте шаблон и сверьтесь с ним")
    body = raw[1:]
    if len(body) > MAX_ROWS:
        raise FileError(f"В файле {len(body)} строк, максимум {MAX_ROWS}")
    first_seen: dict[str, int] = {}
    rows, servers = [], OrderedDict()
    for line, r in body:
        cell = lambda k: str(r[cols[k]]).strip() if cols[k] < len(r) and r[cols[k]] is not None else ""
        name, model_raw = cell("server"), cell("model")
        serial = re.sub(r"\.0$", "", cell("serial"))
        d = _parse_date(r[cols["date"]]) if cols["date"] < len(r) else None
        model = normalize_model(model_raw)
        err = None
        if not name:
            err = "Не указан сервер"
        elif not model:
            err = "Модель не поддерживается: только H100, H200, A100, L40S"
        elif not serial:
            err = "Нет серийного номера"
        elif not SERIAL_RE.match(serial):
            err = "Серийный номер: 13 цифр"
        elif serial in first_seen:
            err = f"Серийный номер повторяется, см. строку {first_seen[serial]}"
        elif d is None:
            err = "Дата ввода: ДД.ММ.ГГГГ"
        elif d > today:
            err = "Дата ввода позже сегодняшней"
        elif name in servers and servers[name]["model"] != model:
            err = f"В сервере {name} смешаны модели {servers[name]['model']} и {model}"
        if serial and serial not in first_seen:
            first_seen[serial] = line
        row = dict(line=line, server=name, model=model_raw, serial=serial, date=d.isoformat() if d else None, ok=err is None, error=err)
        rows.append(row)
        if err is None:
            srv = servers.setdefault(name, dict(name=name, model=model, gpus=[]))
            srv["gpus"].append(dict(idx=len(srv["gpus"]), serial=serial, commissioned=d.isoformat()))
    bad = sum(1 for r in rows if not r["ok"])
    good = [s for s in servers.values()]
    return dict(rows=rows, servers=good, errors=bad, total=len(rows), cards=sum(len(s["gpus"]) for s in good),
                count_servers=len(good), template_hint=None)


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Оборудование"
    ws.append(["Сервер", "Модель", "Серийный номер", "Дата ввода"])
    for c, w in zip("ABCD", (14, 22, 20, 14)):
        ws.column_dimensions[c].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_smi(filename: str, data: bytes) -> list[dict]:
    """CSV из `nvidia-smi --query-gpu=index,serial,name,temperature.gpu,utilization.gpu --format=csv`."""
    rows = [r for r in _rows(filename, data) if any(str(c or "").strip() for c in r)]
    if not rows:
        raise FileError("Файл пуст")
    head = [str(h).strip().lower() for h in rows[0]]

    def find(*keys):
        for i, h in enumerate(head):
            if any(k in h for k in keys):
                return i
        return None
    i_serial, i_temp, i_util = find("serial"), find("temperature", "temp"), find("utilization", "util", "загруз")
    i_host, i_idx, i_hours = find("host", "server", "сервер"), find("index"), find("hours", "наработ")
    if i_serial is None or i_temp is None:
        raise FileError("В файле нет колонок serial и temperature.gpu "
                        "(команда: nvidia-smi --query-gpu=index,serial,name,temperature.gpu,utilization.gpu --format=csv)")
    num = lambda v: float(re.sub(r"[^\d.]", "", str(v)) or 0)
    out = []
    for r in rows[1:]:
        if len(r) <= max(i_serial, i_temp):
            continue
        out.append(dict(
            serial=re.sub(r"\.0$", "", str(r[i_serial]).strip()), temp=num(r[i_temp]),
            util=num(r[i_util]) if i_util is not None and i_util < len(r) else 0,
            host=str(r[i_host]).strip() if i_host is not None and i_host < len(r) else None,
            index=int(num(r[i_idx])) if i_idx is not None and i_idx < len(r) else None,
            hours=num(r[i_hours]) if i_hours is not None and i_hours < len(r) else None,
        ))
    return out


# ---------- проверка реквизитов заёмщика ----------

def inn_valid(inn: str) -> bool:
    """ИНН: 10 (юрлицо) или 12 (ИП/физлицо) цифр с проверкой контрольной суммы."""
    if not re.fullmatch(r"\d{10}|\d{12}", inn or ""):
        return False
    d = [int(c) for c in inn]
    def chk(w):
        return sum(a * b for a, b in zip(w, d)) % 11 % 10
    if len(inn) == 10:
        return chk([2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[9]
    return chk([7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[10] and chk([3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8]) == d[11]


def make_inn(base9: str) -> str:
    """Достраивает контрольную цифру ИНН юрлица (для демо-данных)."""
    d = [int(c) for c in base9]
    return base9 + str(sum(a * b for a, b in zip([2, 4, 10, 3, 5, 9, 4, 6, 8], d)) % 11 % 10)


APP_RE = re.compile(r"^(КД|KD)-\d{4}-\d{4}$", re.I)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def validate_borrower(b: dict) -> dict:
    errs = {}
    name = str(b.get("name", "")).strip()
    if not 2 <= len(name) <= 200:
        errs["name"] = "От 2 до 200 символов"
    if not inn_valid(str(b.get("inn", "")).strip()):
        errs["inn"] = "10 или 12 цифр, проверьте контрольную сумму"
    if not APP_RE.match(str(b.get("application", "")).strip()):
        errs["application"] = "Формат КД-ГГГГ-NNNN, например КД-2026-0418"
    if not str(b.get("location", "")).strip():
        errs["location"] = "Укажите, где стоит оборудование"
    if not str(b.get("contact", "")).strip():
        errs["contact"] = "Укажите контакт в ЦОД"
    em = str(b.get("contact_email", "")).strip()
    if em and not EMAIL_RE.match(em):
        errs["contact_email"] = "Проверьте адрес почты"
    return errs
